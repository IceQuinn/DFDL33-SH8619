from collections import Counter
from hashlib import sha256
from pathlib import Path
import zipfile

from docx import Document

root = Path(r"E:\Project\SVN\DFDL33-SH8619")
target = root / "doc" / "5、开发文档" / "协议转换单元软件开发设计 - 副本.docx"
reference = root / "doc" / "5、开发文档" / "协议转换单元软件开发设计.docx"
doc = Document(target)
headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
with zipfile.ZipFile(target) as package:
    xml = package.read("word/document.xml").decode("utf-8")

print("size", target.stat().st_size)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "sections", len(doc.sections))
print("heading_counts", Counter(style for style, _ in headings))
print("first_headings", headings[:8])
print("toc_field", 'TOC \\o "1-4"' in xml)
print("replacement_chars", xml.count("\ufffd"))
print("reference_sha256", sha256(reference.read_bytes()).hexdigest().upper())
