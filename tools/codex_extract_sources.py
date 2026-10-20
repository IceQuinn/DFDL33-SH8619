from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def extract_docx(path: Path, output: Path) -> None:
    doc = Document(path)
    lines = [f"FILE\t{path}", f"SECTIONS\t{len(doc.sections)}"]
    for i, section in enumerate(doc.sections, 1):
        lines.append(
            "SECTION\t{}\t{}x{}\tMARGINS={} {} {} {}\tORIENTATION={}".format(
                i,
                section.page_width,
                section.page_height,
                section.top_margin,
                section.right_margin,
                section.bottom_margin,
                section.left_margin,
                section.orientation,
            )
        )
    lines.append("BODY")
    for idx, para in enumerate(doc.paragraphs):
        text = clean(para.text)
        if text:
            lines.append(f"P\t{idx}\t{para.style.name}\t{text}")
    for tidx, table in enumerate(doc.tables):
        lines.append(f"TABLE\t{tidx}\tROWS={len(table.rows)}\tCOLS={len(table.columns)}")
        for ridx, row in enumerate(table.rows):
            vals = [clean(cell.text) for cell in row.cells]
            lines.append(f"TR\t{tidx}\t{ridx}\t" + "\t".join(vals))
    for sidx, section in enumerate(doc.sections):
        for label, container in (("HEADER", section.header), ("FOOTER", section.footer)):
            for pidx, para in enumerate(container.paragraphs):
                text = clean(para.text)
                if text:
                    lines.append(f"{label}\t{sidx}\t{pidx}\t{para.style.name}\t{text}")
            for tidx, table in enumerate(container.tables):
                for ridx, row in enumerate(table.rows):
                    vals = [clean(cell.text) for cell in row.cells]
                    lines.append(f"{label}_TR\t{sidx}\t{tidx}\t{ridx}\t" + "\t".join(vals))
    output.write_text("\n".join(lines), encoding="utf-8")


def extract_xlsx(path: Path, output: Path) -> None:
    wb = load_workbook(path, read_only=True, data_only=False)
    result = {"file": str(path), "sheets": []}
    for ws in wb.worksheets:
        rows = []
        for ridx, row in enumerate(ws.iter_rows(values_only=True), 1):
            values = [clean(v) for v in row]
            if any(values):
                while values and values[-1] == "":
                    values.pop()
                rows.append({"row": ridx, "values": values})
        result["sheets"].append(
            {"title": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "rows": rows}
        )
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("kind", choices=("docx", "xlsx"))
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    if args.kind == "docx":
        extract_docx(args.input, args.output)
    else:
        extract_xlsx(args.input, args.output)


if __name__ == "__main__":
    main()
