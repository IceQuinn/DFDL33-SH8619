from pathlib import Path
from copy import deepcopy
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\Project\SVN\DFDL33-SH8619")
TARGET = ROOT / "doc" / "5、开发文档" / "协议转换单元软件开发设计 - 副本.docx"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-4" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在打开文档时更新"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, placeholder, fld_end])


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, txt, end])


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def body_p(doc, text="", bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.74)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    p.add_run("• ")
    p.add_run(text)
    return p


def heading(doc, text, level, page_break=False):
    if page_break:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
    return doc.add_heading(text, level=level)


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(h)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        if widths:
            set_cell_width(cell, widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(font_size)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cell, widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def status_table(doc, rows, font_size=9):
    return add_table(doc, ["功能项", "规范/点表要求", "当前软件实现", "状态与说明"], rows,
                     widths=[3.0, 5.0, 5.3, 3.8], font_size=font_size)


def build():
    doc = Document(TARGET)
    clear_body(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)

    styles = doc.styles
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = styles[name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(12)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["Heading 2"].font.size = Pt(18)
    styles["Heading 3"].font.size = Pt(16)
    styles["Heading 4"].font.size = Pt(14)

    title = doc.add_heading("协议转换单元软件功能设计", level=1)
    title.paragraph_format.space_after = Pt(18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("装置型号：DFDL33-SH8619")
    r.font.size = Pt(14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("依据技术规范书、DL/T 645点表及当前软件实现整理")
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("文档版本：V1.0    日期：2026-10-20")

    doc.add_page_break()
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = toc_title.add_run("目录")
    rt.bold = True
    rt.font.size = Pt(18)
    toc = doc.add_paragraph()
    add_toc(toc)
    doc.add_page_break()

    heading(doc, "1 文档说明", 2)
    heading(doc, "1.1 编制目的", 3)
    body_p(doc, "本文从技术规范、DL/T 645点表、软硬件接口、逆变器协议库、应用程序和Boot程序六个维度，对DFDL33-SH8619低压分布式光伏协议转换单元进行功能分解。内容重点回答装置有哪些大功能、每个大功能包含哪些小功能和处理细节、功能之间如何关联，以及当前代码与规范要求之间的边界。")
    heading(doc, "1.2 依据与优先级", 3)
    add_table(doc, ["依据", "主要用途", "冲突时处理原则"], [
        ["0923版技术规范书", "确定装置目标能力、接口、性能、事件、升级和安全要求", "作为需求基线"],
        ["DLT64点表", "确定645参数、动作、档案和协议库数据标识及数据结构", "作为645点位细化依据"],
        ["软硬件接口文档", "确认串口、指示灯、采样和Flash等物理资源", "作为硬件映射依据"],
        ["逆变器通用协议库", "确认厂家协议、寄存器、类型、字节序和控制能力", "作为下行适配依据"],
        ["当前APP/Boot源码", "确认实际运行流程、线程、缓存、持久化和异常行为", "作为“当前已实现”判定依据"],
        ["研发自测用例", "交叉检查场景、边界和预期交互", "用于发现差异，不替代源码事实"],
    ], widths=[4.2, 7.1, 5.8])
    heading(doc, "1.3 状态说明", 3)
    bullet(doc, "已实现：当前代码中存在完整入口、处理过程及输出/持久化链路。")
    bullet(doc, "部分实现：主体框架存在，但接口、存储、事件挂接、安全校验或某一路通道尚未闭环。")
    bullet(doc, "规范要求：规范或点表提出，但在当前代码中未找到完整实现。")
    body_p(doc, "本文描述的是当前软件基线与需求目标的结合视图。涉及“规范要求”或“待闭环”的内容，不应直接作为已通过验证的功能承诺。")

    heading(doc, "2 装置总体功能框架", 2, page_break=True)
    heading(doc, "2.1 总体定位", 3)
    body_p(doc, "装置位于主站/集中器与分布式光伏逆变器之间：上行以DL/T 645及其扩展数据标识向主站提供统一接口，下行以Modbus主站方式适配不同厂家逆变器；同时完成档案和协议库管理、实时采集、柔性调节、事件记录、本地维护、电压采集和软件升级。")
    heading(doc, "2.2 大功能分层", 3)
    add_table(doc, ["层级", "大功能", "核心职责", "主要关联"], [
        ["接入层", "上行通信", "RS485-I、载波、蓝牙接收DL/T 645请求并返回响应", "645服务、参数、升级、指示灯"],
        ["接入层", "下行通信", "RS485-II及两组RJ45共五路有线接口，Modbus RTU主站收发", "识别、采集、控制、转接"],
        ["协议层", "DL/T 645服务", "报文校验、地址匹配、数据标识路由、读写/控制应答", "全部上行业务"],
        ["协议层", "逆变器协议库", "用厂家、版本、特征寄存器和点位描述屏蔽厂家差异", "识别、采集、控制、档案"],
        ["业务层", "自适应识别与档案", "扫描地址、匹配协议、建立最多12台逆变器档案", "协议库、周期采集、控制"],
        ["业务层", "数据采集与转换", "周期/实时读取、类型解析、单位缩放、645编码和无效值处理", "缓存、主站、运行状态"],
        ["业务层", "柔性调节", "启停、有功/无功/功率因数、百分比及广播控制", "档案、协议库、事件"],
        ["业务层", "时段控制", "每台设备两时段计划、执行、恢复和失败重试", "RTC、控制队列、档案"],
        ["管理层", "参数/维护/事件/存储", "配置、时钟、位置、校准、日志和A/B冗余保存", "645、本地维护、可靠性"],
        ["生命周期", "软件升级与Boot", "分包接收、断点续传、CRC校验、搬运和跳转", "上行通道、外部Flash、看门狗"],
    ], widths=[2.0, 3.1, 6.6, 5.4], font_size=9)
    heading(doc, "2.3 核心数据流", 3)
    heading(doc, "2.3.1 主站读数据链路", 4)
    body_p(doc, "主站请求经RS485-I、载波或蓝牙进入统一DL/T 645接收队列；645线程完成地址、控制码、长度和校验和检查，再按数据标识进入点表。装置参数直接从配置或RTC读取；逆变器数据可触发实时刷新或读取周期缓存，随后转换为BCD/SBCD/ASCII/原始结构并通过原通道返回。")
    heading(doc, "2.3.2 主站写参数/控制链路", 4)
    body_p(doc, "写请求经点表校验长度、范围和权限。装置参数写入配置并通过A/B区保存；逆变器控制根据档案定位端口、根据协议库形成Modbus 06/10报文、进入对应端口异步队列，等待成功、异常或超时结果后形成645写应答。")
    heading(doc, "2.3.3 自动识别与运行链路", 4)
    body_p(doc, "空闲主端口扫描Modbus地址1～10，读特征寄存器并与协议库范围匹配；识别成功后保存档案。周期采集按档案和协议库读取运行数据、更新有效性缓存；控制和时段控制复用同一端口串行队列，避免同一总线并发冲突。")

    heading(doc, "3 软件运行框架", 2, page_break=True)
    heading(doc, "3.1 启动初始化顺序", 3)
    add_table(doc, ["顺序", "初始化项", "作用与依赖"], [
        ["1", "系统运行时间", "记录上电时刻，为日志和运行状态提供基准"],
        ["2", "装置信息输出", "打印型号、版本等启动信息"],
        ["3", "外部Flash", "为配置、协议库、档案、事件、时段计划和升级区提供访问基础"],
        ["4", "装置配置", "A/B区校验装载；失败或版本不匹配时恢复默认"],
        ["5", "DL/T 645", "建立地址、数据点和报文处理环境"],
        ["6", "逆变器协议库", "先装载默认/持久化协议，供档案绑定"],
        ["7", "逆变器档案", "装载最多12台档案并建立协议指针"],
        ["8", "时段控制运行态", "在线程启动前恢复计划及上下文"],
        ["9", "事件", "装载两类事件区、创建事件队列并记录上电事件"],
        ["10", "串口与蓝牙", "配置七个UART逻辑通道及HJ-02C模块"],
        ["11", "业务线程", "资源就绪后统一创建九个应用线程"],
    ], widths=[1.5, 4.5, 10.9])
    heading(doc, "3.2 线程与并发模型", 3)
    add_table(doc, ["线程", "优先级/栈", "主要职责", "同步方式"], [
        ["hj02c_rx", "14 / 2048", "接收蓝牙SPI数据并送入645入口", "SPI + 645队列"],
        ["voltage_acq", "15 / 1024", "ADC/DMA采样、RMS计算和校准", "DMA信号量"],
        ["uart_mgmt", "17 / 1024", "串口静默分帧、协议分发、转接协调", "中断缓存 + 轮询"],
        ["cycle", "17 / 4096", "自动识别、周期抄读、实时刷新和控制队列", "端口状态机/消息队列"],
        ["upgrade", "18 / 6144", "升级摘要、分包、校验和应答", "升级消息队列"],
        ["645_sl", "20 / 2048", "统一处理串口/蓝牙645报文", "环形队列 + 信号量"],
        ["time_ctrl", "22 / 2048", "两时段计划执行、恢复和重试", "200 ms循环 + 控制队列"],
        ["led_run", "25 / 512", "运行灯和收发灯模式/超时", "125 ms节拍"],
        ["event", "30 / 1024", "事件入队、环形保存和A/B落盘", "事件消息队列"],
    ], widths=[2.5, 2.6, 7.2, 4.6], font_size=9)
    heading(doc, "3.3 串行化原则", 3)
    body_p(doc, "同一物理Modbus总线上的自动识别、周期抄读、实时读、柔性控制和通信转接必须串行执行。软件通过三组主端口状态机和控制队列统一调度；第二路RJ45转接报文由对应主端口代发，响应需按地址、功能码和CRC匹配后返回。")

    heading(doc, "4 通信接口与通道管理", 2, page_break=True)
    heading(doc, "4.1 物理/逻辑通道映射", 3)
    add_table(doc, ["逻辑UART", "外部接口", "默认角色", "默认参数", "业务"], [
        ["UART3", "RS485-I", "DL/T 645从机", "9600, 8E1", "上行/维护"],
        ["UART8", "载波", "DL/T 645从机", "9600, 8E1", "上行主站"],
        ["UART6", "RS485-II", "Modbus主站", "9600, 8N1", "主下行端口1"],
        ["UART4", "RJ45一路", "Modbus主站", "9600, 8N1", "主下行端口2"],
        ["UART7", "RJ45一路", "Modbus主站", "9600, 8N1", "主下行端口3"],
        ["UART1", "RJ45第二路", "转接从入口", "9600, 8N1", "与UART4成对"],
        ["UART5", "RJ45第二路", "转接从入口", "9600, 8N1", "与UART7成对"],
        ["SPI4", "HJ-02C蓝牙", "DL/T 645透传", "模块命令接口", "本地维护"],
    ], widths=[2.4, 3.2, 3.2, 3.0, 5.1], font_size=9)
    heading(doc, "4.2 串口接收与分帧", 3)
    bullet(doc, "各UART按字节中断接收，单通道静态缓存512字节。")
    bullet(doc, "连续10 ms无新字节判定一帧结束；管理线程约1 ms轮询各通道。")
    bullet(doc, "UART3/UART8送DL/T 645；UART6/UART4/UART7送识别、抄读或控制；UART1/UART5参与成对转接。")
    bullet(doc, "收发动作触发对应上行、下行或维护指示灯，并在超时后自动熄灭。")
    heading(doc, "4.3 通信参数配置", 3)
    body_p(doc, "点表04000800～04000806分别配置RS485-I、RS485-II、两组RJ45四路串口及载波，数据为4字节小端波特率加1字节校验格式。当前代码允许2400、4800、9600、19200、38400，写入后保存但需重启才能按新参数重新初始化；规范允许最高115200，点表中的04000807无线参数在当前645点表实现中明确未支持。")

    heading(doc, "5 上行DL/T 645服务", 2, page_break=True)
    heading(doc, "5.1 报文与地址处理", 3)
    bullet(doc, "校验帧头68H、第二帧头、数据长度、校验和与16H结束符；数据域按规约执行减33H/加33H。")
    bullet(doc, "支持本机单播地址、广播地址和AA/99通配地址；支持读数据11H、写数据14H、读地址13H、写地址15H、广播校时08H及私有升级控制。")
    bullet(doc, "写地址执行BCD合法性检查，保存装置地址，并同步更新蓝牙广播名称。")
    bullet(doc, "RS485-I、载波和蓝牙共享同一645处理线程，普通响应按请求来源返回原通道。")
    heading(doc, "5.2 数据点路由", 3)
    body_p(doc, "数据中心以数据标识、掩码、访问权限、编解码方式、选择器范围、数据长度、比例和上下限为元数据，统一调用读写回调。DI0常用于选择逆变器01～0C、协议槽位或FF全选。非法数据标识、长度、权限、范围或业务失败由645顶层转为异常应答。")
    heading(doc, "5.3 645电表兼容点", 3)
    body_p(doc, "为兼容主站按电表方式访问，软件实现了部分标准电能量、瞬时量、冻结/曲线点。该装置不是计量表，除本机单相电压外，大量电流、功率、功率因数和电能点采用零值、FF无效相或空记录应答，避免主站因“不支持数据标识”中断流程。")
    heading(doc, "5.3.1 电能量与冻结点", 4)
    add_table(doc, ["数据标识/范围", "含义", "当前应答"], [
        ["00010000～00010400、0001FF00", "正向有功总/费率及组合", "固定零值"],
        ["00020000～00020400、0002FF00", "反向有功总/费率及组合", "固定零值"],
        ["05060101、05060201", "上一次日冻结正/反向有功", "固定零值"],
        ["06000001", "给定时刻负荷记录", "正常空应答，仅回显DI"],
    ], widths=[5.0, 6.3, 5.6], font_size=9)
    heading(doc, "5.3.2 瞬时量与曲线点", 4)
    add_table(doc, ["点组", "A相/总量", "B、C相", "细节"], [
        ["02010100、0201FF00；06100101、061001FF", "返回ADC采集并校准的本机A相电压，0.1 V", "FF", "瞬时量与曲线点共用电压源"],
        ["0202/061002电流组", "固定0", "FF", "未配置电流采样"],
        ["0203/0204/0205及061003/061004功率组", "总量/A相固定0", "FF", "含有功、无功、视在功"],
        ["0206及061005功率因数组", "总量/A相固定0", "FF", "格式X.XXX"],
        ["061006、061007电能曲线组", "固定0", "固定0", "用于兼容主站曲线读取"],
    ], widths=[4.5, 4.8, 2.2, 5.4], font_size=9)
    heading(doc, "5.4 装置参数与动作点", 3)
    add_table(doc, ["数据标识", "功能", "读写", "数据/行为细节"], [
        ["0400040F", "经纬度和高度", "读/写", "经度4字节BCD、纬度4字节BCD、高度3字节BCD；写后保存"],
        ["04000800～04000806", "串口参数", "读/写", "4字节小端波特率+1字节校验；重启生效"],
        ["04000900", "周期抄读时间", "读/写", "2字节BCD；0关闭，5～3600 s"],
        ["04000A00", "装置重启", "写", "写01触发重启"],
        ["04000B00", "清除事件", "写", "写01清除全部事件分类"],
        ["04000C00", "恢复默认", "写", "写01恢复并保存默认配置，不自动重启"],
        ["04000D00", "电压校准", "读/写", "读回4字节浮点系数；写01触发基于220.0 V的异步校准"],
        ["04800001", "软件版本", "读", "固定32字节ASCII，零填充"],
    ], widths=[3.2, 3.5, 2.0, 8.2], font_size=9)
    heading(doc, "5.5 逆变器扩展点", 3)
    add_table(doc, ["点组", "功能", "选择范围", "刷新/编码细节"], [
        ["02E601～02E605", "电压、电流、有功、无功、功率因数", "DI0=01～0C或FF", "读取前可实时刷新；失败点返回FF；有功/无功采用带符号BCD"],
        ["02E60F", "单台全部变量55字节", "DI0=01～0C", "按电压、电流、有功、无功、功率因数顺序拼接，不支持FF聚合"],
        ["04E601～04E603", "额定P/Q、输出类型", "DI0=01～0C或FF", "来自协议库定义和设备实时缓存"],
        ["04E604～04E609", "启停及数值/百分比控制", "DI0=01～0C或FF", "读回缓存；写入异步Modbus控制，FF逐台执行"],
        ["04E60A", "日发电量", "DI0=01～0C或FF", "格式XXXXXX.XX kWh"],
        ["04E60B、04E60C", "有功数值/百分比时段控制", "DI0=01～0C或FF", "每台最多2时段，当天执行"],
        ["04E62100/01～0C", "档案数量/档案", "数量固定点；档案按DI0", "每档案36字节"],
        ["04E701FE/01～64", "协议有效数/协议槽位", "有效数固定点；槽位按DI0", "单协议238字节；代码容量100槽，但645选择器仅1～100，注释“01～64”为十六进制表示"],
    ], widths=[3.8, 4.5, 3.1, 5.5], font_size=8.8)
    heading(doc, "5.6 实时刷新与无效值策略", 3)
    body_p(doc, "逆变器实时读在开始时先将目标缓存点置无效，再按档案顺序刷新，整批最长约55 s，为主站60 s超时预留组帧和发送余量。尚未执行、超时或解析失败的点保持无效，645编码为FF；内部可保留旧原始值用于诊断，但不会伪装成新鲜有效数据。")

    heading(doc, "6 下行Modbus与通信转接", 2, page_break=True)
    heading(doc, "6.1 Modbus主站能力", 3)
    bullet(doc, "读支持03H/04H；写支持06H/10H；校验地址、功能码、长度、CRC及异常码。")
    bullet(doc, "按协议库描述解析有符号/无符号8/16/32位、Float32/Float64、ASCII、BCD、BCD_TIME和位域。")
    bullet(doc, "支持ABCD、CDAB、BADC、DCBA字节序及小数位缩放，输出统一工程值。")
    bullet(doc, "每个主下行端口使用独立非阻塞状态机，超时或异常不会阻塞其他端口。")
    heading(doc, "6.2 通信转接", 3)
    body_p(doc, "两路RJ45的第二串口作为外部转接入口，分别借用对应第一串口向逆变器发送。转接请求在自动识别、采集和控制之间排队；响应依据从站地址、功能码和CRC与在途请求匹配后返回原入口。新的待处理报文可覆盖旧待处理报文，并在必要时接替旧在途响应，属于单缓存式转接。")
    heading(doc, "6.3 链路监测", 3)
    body_p(doc, "当前代码通过请求超时、Modbus异常、CRC/报文不匹配和缓存有效性体现链路状态，能够阻止错误数据上送。规范要求的独立下行链路状态统计、通信异常事件及主动上报，在当前事件点和上送链路中尚未形成完整闭环。")

    heading(doc, "7 逆变器协议库与自动识别", 2, page_break=True)
    heading(doc, "7.1 协议库结构", 3)
    body_p(doc, "协议库以厂家ASCII名称、规约版本、特征寄存器、数据点、参数点和控制点组成。单条线格式238字节；代码容量100条。每个寄存器块包含地址、寄存器数、功能码、数据类型、字节序和小数位，使识别、采集和控制共享同一描述。")
    heading(doc, "7.2 默认协议覆盖", 3)
    body_p(doc, "固件默认库包含阳光、华为V2/V3、固德威、古瑞瓦特、三晶、锦浪、首航、上能、奥克斯、特变电工、爱士惟、禾望、银河耐吉等多组协议版本。并非每个协议都具备全部变量和控制点，具体能力由寄存器块有效性决定。")
    heading(doc, "7.3 自动识别流程", 3)
    heading(doc, "7.3.1 地址扫描", 4)
    body_p(doc, "三个主有线端口分别扫描从站地址1～10，使用03H读取寄存器0、数量1作为在线探测，单次约1 s超时。已有有效档案占用的端口不重复扫描；空端口完成扫描后转入周期抄读。")
    heading(doc, "7.3.2 协议匹配", 4)
    body_p(doc, "探测到设备后遍历有效协议，按每条协议配置的特征寄存器读取并解码，将报文值与特征下限/上限比较。匹配成功即形成地址、端口、厂家和版本档案并持久化。")
    heading(doc, "7.3.3 识别边界", 4)
    body_p(doc, "当前自动识别只覆盖三个主有线端口和地址1～10；规范中的无线接口可接最多12台逆变器，但当前代码未见无线端口扫描和无线档案建链。识别结果总容量受12个全局档案槽限制。")
    heading(doc, "7.4 协议库维护", 3)
    body_p(doc, "启动时协议库可由Flash A/B区装载并由默认库补充。645可读无效槽位为238字节FF，并可校验写入RAM后立即供识别/采集使用；当前645写协议路径明确不执行Flash保存，因此重启后临时写入可能丢失，需要在正式维护流程中补齐持久化或明确仅用于调试。")

    heading(doc, "8 档案管理与数据采集", 2, page_break=True)
    heading(doc, "8.1 档案管理", 3)
    body_p(doc, "装置最多管理12台逆变器。每个档案固定36字节：1字节Modbus地址、32字节厂家ASCII、2字节规约版本、1字节端口号。档案通过A/B区保存；加载时按厂家与版本绑定协议指针。协议不存在时档案仍可保留，但标记为无效，不参与采集或控制。")
    bullet(doc, "支持档案数量读取、单档案读取/写入、增加、设置和删除。")
    bullet(doc, "档案变化会重新建立协议指针，并与时段控制协调，防止旧档案计划误控新设备。")
    bullet(doc, "端口字段把档案与三组主下行总线关联；转接副口不单独形成采集档案。")
    heading(doc, "8.2 周期抄读", 3)
    body_p(doc, "全局周期由04000900配置，0表示关闭，5～3600 s有效。周期线程按档案和端口排队读取变量/参数；相邻相电压、电流等连续寄存器在条件允许时合并请求。解析成功更新值与有效标志，失败只清除有效标志。")
    heading(doc, "8.3 实时缓存", 3)
    add_table(doc, ["缓存内容", "来源", "更新触发", "失效条件"], [
        ["三相电压/电流", "逆变器寄存器", "周期或645实时读", "超时、异常、解析失败"],
        ["总/分相有功和无功", "逆变器寄存器", "周期或645实时读", "同上"],
        ["总/分相功率因数", "逆变器寄存器", "周期或645实时读", "同上"],
        ["额定P/Q、输出类型、日发电量", "协议点/逆变器寄存器", "周期或实时读", "协议不支持或通信失败"],
        ["运行状态", "总有功推断", "约1 s判断", "无有效有功数据时为未知"],
        ["控制设定值", "最近控制/读取结果", "控制执行或读取", "档案/协议变化"],
    ], widths=[4.1, 4.2, 4.4, 4.2], font_size=9)
    heading(doc, "8.4 数据转换", 3)
    body_p(doc, "协议库先把Modbus寄存器解码为统一整数/浮点工程值，DL/T 645层再按点表要求转换为BCD、带符号BCD、ASCII或原始结构。符号、比例、小数位、字节序和无效FF处理分层完成，避免不同厂家格式渗透到上行接口。")

    heading(doc, "9 柔性调节与时段控制", 2, page_break=True)
    heading(doc, "9.1 控制类型", 3)
    add_table(doc, ["控制", "645点", "范围/格式", "下行实现"], [
        ["开机/关机", "04E604", "0/1", "协议库固定控制值，Modbus 06/10"],
        ["有功数值", "04E605", "带符号XXXX.XXXX kW", "按协议数据类型/倍率写入"],
        ["无功数值", "04E606", "带符号XXXX.XXXX kvar", "同上"],
        ["功率因数", "04E607", "-1.000～1.000", "同上"],
        ["有功百分比", "04E608", "0.0%～110.0%", "同上"],
        ["无功百分比", "04E609", "-110.0%～110.0%", "同上"],
        ["有功数值时段", "04E60B", "每天最多2时段", "计划线程在起点写目标，终点恢复"],
        ["有功百分比时段", "04E60C", "每天最多2时段", "终点默认恢复100%"],
    ], widths=[3.3, 2.4, 5.0, 6.2], font_size=9)
    heading(doc, "9.2 单台与广播控制", 3)
    body_p(doc, "DI0=01～0C定位单台档案，FF按12个固定槽位逐台执行。执行前检查档案有效、协议支持、数值范围和端口可用；同端口请求按队列串行。结果区分成功、不支持、Modbus异常和超时，再映射为645正常或异常写应答。")
    heading(doc, "9.3 时段计划细节", 3)
    bullet(doc, "每台逆变器每天最多2个时段；检查hhmm有效、起止顺序、时段不重叠且不跨日。")
    bullet(doc, "线程约200 ms检查计划；开始时下发目标，结束时恢复有功百分比100%。")
    bullet(doc, "单次控制结果等待约3 s；恢复失败后每60 s重试。")
    bullet(doc, "计划按当天语义执行，跨日清理失效计划；档案替换或清除会停止/清理相关计划。")
    bullet(doc, "当前仅支持有功数值和有功百分比时段控制，不含无功时段控制。")
    heading(doc, "9.4 控制关联", 3)
    body_p(doc, "控制依赖档案确定地址/端口，依赖协议库确定寄存器和编码，依赖UART/Modbus完成发送，依赖缓存提供读回；时段控制还依赖RTC和持久化计划。规范要求将启停、有功、无功、功率因数控制形成独立事件记录，当前事件框架尚未挂接这些控制结果。")

    heading(doc, "10 本地维护、时钟与参数", 2, page_break=True)
    heading(doc, "10.1 蓝牙维护", 3)
    body_p(doc, "HJ-02C通过SPI4连接。启动时软件对模块执行恢复出厂、读取版本并把装置645地址设置为广播名称；蓝牙收到的数据直接送入统一645处理链路，普通应答经SPI返回。模块接口还具备广播开关/间隔、连接间隔、发射功率、高速模式、复位和断开连接命令。")
    heading(doc, "10.2 时钟", 3)
    bullet(doc, "支持RTC读取和广播校时08H；时段控制和事件时间戳均依赖RTC。")
    bullet(doc, "规范要求24 h误差不超过0.5 s、对时误差小于10 s，并记录最近50次校时前后时间。当前代码具备RTC设置入口，但未见完整精度校验与校时事件闭环。")
    heading(doc, "10.3 参数和位置", 3)
    body_p(doc, "装置配置包含7个逻辑UART的协议/波特率/校验、645地址、经度、纬度、高度、电压校准系数及全局抄读周期。配置使用版本号与CRC的A/B冗余保存。Flash校验失败或版本不匹配时恢复默认；一般恢复默认保留645地址，完全恢复默认同时恢复地址。")
    heading(doc, "10.4 本地安全边界", 3)
    body_p(doc, "技术规范要求本地维护具备身份认证、权限验证，升级也需鉴权。当前蓝牙接收路径直接复用645协议，未见独立用户认证、会话权限或报文签名逻辑；且启动时恢复蓝牙模块出厂设置可能覆盖模块侧安全配置。该项应作为上线前安全闭环重点。")

    heading(doc, "11 单相电压采集与状态指示", 2, page_break=True)
    heading(doc, "11.1 电压采集", 3)
    body_p(doc, "ADC1由定时器触发，DMA完成一批采样后释放信号量；采样线程计算RMS，将连续5次结果平均后乘以配置中的校准系数，得到0.1 V单位的本机电压。每轮处理后空闲约500 ms再启动下一轮。该值直接供02010100/0201FF00及对应曲线电压点使用。")
    heading(doc, "11.2 电压校准", 3)
    body_p(doc, "写04000D00=01只设置校准标志。下一次采样处理中取最多50次RMS缓存平均值，以220.0 V为基准重新计算系数并保存。读04000D00返回当前4字节浮点校准系数。校准依赖现场电压确为220.0 V，维护流程应明确校准条件。")
    heading(doc, "11.3 指示灯", 3)
    add_table(doc, ["指示灯", "作用", "软件模式"], [
        ["RUN", "装置运行状态", "默认普通闪烁；支持常亮、慢闪、快闪、关闭"],
        ["维护发送/接收", "蓝牙/维护通信活动", "通信触发，超时自动熄灭"],
        ["上行发送/接收", "RS485-I/载波645活动", "通信触发，超时自动熄灭"],
        ["下行发送/接收", "Modbus/转接活动", "通信触发，超时自动熄灭"],
        ["全灯模式", "测试/告警提示", "可覆盖所有灯，超时后恢复原模式"],
    ], widths=[3.8, 7.0, 6.1])

    heading(doc, "12 事件记录与数据存储", 2, page_break=True)
    heading(doc, "12.1 事件框架", 3)
    body_p(doc, "事件模块分为自检类和操作类，每类最多256条，采用环形索引保存。业务通过消息队列提交事件，事件线程补充时间戳、动作/复归、持续时间和动作值，随后按事件分类写入外部Flash A/B区。支持查询、清除、前后条目定位和掉电恢复。")
    heading(doc, "12.2 当前已接入事件", 3)
    add_table(doc, ["事件", "分类", "触发情况"], [
        ["装置重启", "操作类", "事件描述已定义，具体触发需由系统调用"],
        ["装置失电", "操作类", "事件描述已定义"],
        ["装置上电", "操作类", "Event_Init后主动入队"],
        ["超级电容故障", "自检类", "事件描述已定义，依赖检测逻辑触发"],
    ], widths=[4.5, 3.5, 8.9])
    heading(doc, "12.3 规范事件与当前差异", 3)
    status_table(doc, [
        ["电压越限", "累计次数、最近100次，含发生时间和电压", "未见阈值判断和专用事件点接入", "规范要求"],
        ["逆变器启停控制", "最近100次", "控制链存在，未挂接事件", "部分实现"],
        ["有功/无功/PF控制", "各最近100次", "控制链存在，未挂接事件", "部分实现"],
        ["通信异常", "最近50次", "存在超时/异常判定，未挂接事件", "部分实现"],
        ["事件清除", "最近50次", "可清除，但清除动作本身未形成保留事件", "部分实现"],
        ["软件升级", "最近50次，含前后版本", "升级链存在，未挂接事件", "部分实现"],
        ["校时", "最近50次，含前后时间", "校时入口存在，未挂接事件", "部分实现"],
    ])
    heading(doc, "12.4 存储对象与掉电保护", 3)
    add_table(doc, ["对象", "介质/方式", "掉电恢复", "备注"], [
        ["装置配置", "外部Flash A/B + CRC/版本", "支持", "地址、串口、位置、校准、抄读周期"],
        ["逆变器档案", "外部Flash A/B", "支持", "最多12条"],
        ["协议库", "外部Flash A/B + 固件默认库", "支持", "645临时写当前不落盘"],
        ["时段计划", "外部Flash运行配置", "支持", "需与档案变化联动"],
        ["事件", "分区环形 + A/B", "支持", "两类各最多256条"],
        ["实时采集缓存", "RAM", "不支持", "重启后由周期/实时读重建"],
        ["升级包", "外部Flash固定区", "支持", "位图支持断点续传"],
    ], widths=[3.5, 5.3, 3.0, 5.1], font_size=9)

    heading(doc, "13 软件升级与Boot", 2, page_break=True)
    heading(doc, "13.1 APP侧升级", 3)
    heading(doc, "13.1.1 摘要与分包", 4)
    body_p(doc, "私有645升级报文进入升级消息队列。摘要记录升级类型、版本、文件大小、分包总数和文件CRC；数据包写入外部Flash固定升级区，并在200字节位图中标记已接收分包。主站可读取位图定位缺包，实现断点续传和选择性重发。")
    heading(doc, "13.1.2 完整性检查与重启", 4)
    body_p(doc, "全部分包收齐后对外部Flash中的完整镜像计算CRC32。成功则保存带CRC16的升级摘要/标志并触发重启；失败保留缺失或错误状态供重传。升级过程按块操作并配合看门狗。")
    heading(doc, "13.2 Boot侧升级", 3)
    body_p(doc, "Boot读取升级摘要，先校验摘要CRC16和升级标志，再对外部Flash镜像做完整CRC32校验。通过后擦除APP内部Flash区，分块搬运镜像；完成后擦除升级摘要标志并跳转APP。摘要无效、无升级标志或镜像校验失败时直接跳转原APP，防止盲目覆盖。")
    heading(doc, "13.3 升级边界", 3)
    status_table(doc, [
        ["断点续传", "本地/远程升级支持断点续传", "分包位图与缺包应答已实现", "已实现"],
        ["完整性", "升级包校验", "摘要CRC16 + 镜像CRC32 + Boot复验", "已实现"],
        ["来源通道", "应支持本地和远程", "升级应答函数当前固定经HJ-02C蓝牙发送，未保留串口来源", "部分实现"],
        ["鉴权", "需权限和身份认证", "未见升级签名、用户权限或安全会话", "规范要求"],
        ["版本事件", "记录升级前后版本", "未挂接事件模块", "部分实现"],
        ["失败回退", "升级失败不影响运行", "校验失败跳原APP；未见双APP镜像回滚", "部分实现"],
    ])

    heading(doc, "14 可靠性、异常处理与安全", 2, page_break=True)
    heading(doc, "14.1 可靠性机制", 3)
    bullet(doc, "配置、档案、协议库和事件使用CRC及A/B冗余区，单区损坏时可从另一份恢复。")
    bullet(doc, "通信采用CRC、长度、地址、功能码和异常码多层检查；超时会使数据失效而非继续上送旧值。")
    bullet(doc, "三个主下行端口独立状态机，单端口故障不应阻塞其他端口。")
    bullet(doc, "应用和Boot均配置看门狗驱动；升级擦写等长操作主动喂狗。")
    bullet(doc, "Boot在写内部Flash前复验摘要和完整镜像，失败时保留原APP并跳转。")
    heading(doc, "14.2 异常处理", 3)
    add_table(doc, ["异常", "检测", "处理"], [
        ["645帧错误", "帧头、长度、校验和、结束符", "丢弃或异常应答，不进入业务"],
        ["不支持/非法数据标识", "点表掩码和权限查找", "645无请求数据或写失败应答"],
        ["Modbus超时/异常", "端口状态机定时、异常码", "结束本请求、清除有效标志、返回控制失败"],
        ["协议不支持某点", "寄存器块有效性", "读返回FF，控制返回不支持"],
        ["Flash配置损坏", "A/B区CRC和版本", "恢复另一份或恢复默认"],
        ["档案与协议不匹配", "厂家/版本绑定失败", "保留档案但停止采集和控制"],
        ["时段恢复失败", "3 s结果等待", "每60 s重试恢复"],
        ["升级包错误", "位图、CRC16、CRC32", "请求重传或放弃升级并跳原APP"],
    ], widths=[4.0, 6.0, 6.9], font_size=9)
    heading(doc, "14.3 安全缺口", 3)
    body_p(doc, "当前主要安全措施是通信地址、报文格式、CRC和访问权限位，这些属于正确性/完整性检查，不等同于身份认证和授权。规范要求的本地维护认证、升级鉴权、安全初始化及可能的“电鸿”安全能力，需增加用户/密钥管理、挑战应答或签名校验、权限分级、失败锁定和安全审计。")

    heading(doc, "15 功能关联关系", 2, page_break=True)
    heading(doc, "15.1 依赖关系矩阵", 3)
    add_table(doc, ["功能", "直接依赖", "向下影响", "关键失效传播"], [
        ["DL/T 645服务", "上行UART/蓝牙、地址、点表", "参数、采集、控制、档案、升级", "上行不可用时主站所有业务不可达"],
        ["协议库", "Flash、默认库", "识别、采集、控制、档案有效性", "协议缺失导致档案失效或点位不支持"],
        ["自动识别", "下行端口、Modbus、协议库", "档案创建", "识别失败则不能自动采集/控制"],
        ["档案", "识别或645维护、协议库", "采集、控制、时段控制", "地址/端口错误会访问错误设备"],
        ["周期/实时采集", "档案、协议、Modbus", "645上送、运行状态", "通信失败使点无效并返回FF"],
        ["柔性控制", "档案、协议、端口队列", "逆变器运行目标、控制读回", "不支持/超时转为645失败"],
        ["时段控制", "RTC、档案、控制、计划存储", "定时目标与恢复", "时钟或档案变化会影响执行"],
        ["事件", "RTC、业务触发、Flash", "维护查询、故障追溯", "未挂接的业务无法形成审计"],
        ["电压采集", "ADC/DMA、校准参数", "标准645电压点、电压事件", "校准错误直接影响上送电压"],
        ["升级", "645入口、外部Flash、Boot", "软件版本和业务能力", "通道/鉴权/镜像错误会中断升级"],
    ], widths=[3.1, 5.2, 5.0, 3.6], font_size=8.8)
    heading(doc, "15.2 典型业务关联", 3)
    heading(doc, "15.2.1 新逆变器接入", 4)
    body_p(doc, "下行端口发现设备 → 自动识别读取特征 → 协议库匹配 → 档案A/B保存 → 建立协议指针 → 周期采集缓存 → 主站通过02E6/04E6点读取 → 如需控制则复用档案和协议下发。")
    heading(doc, "15.2.2 主站实时读取", 4)
    body_p(doc, "645请求 → 点表选择逆变器 → 清除目标点有效标志 → 端口状态机实时Modbus读取 → 协议解码 → 缓存有效化 → 645编码 → 原上行通道响应；任一步失败则目标字段FF。")
    heading(doc, "15.2.3 主站柔性控制", 4)
    body_p(doc, "645写点 → 选择单台/FF全选 → 档案定位 → 协议确认能力和范围 → 编码Modbus 06/10 → 端口队列串行 → 等待结果 → 645正常/异常应答；规范期望同时生成控制事件。")
    heading(doc, "15.2.4 软件升级", 4)
    body_p(doc, "645私有升级 → APP分包与位图 → 外部Flash镜像 → CRC32 → 保存升级标志并重启 → Boot摘要/镜像复验 → 内部Flash更新 → 清标志 → 跳转新APP → 规范期望记录前后版本事件。")

    heading(doc, "16 规范覆盖与待闭环项", 2, page_break=False)
    heading(doc, "16.1 功能覆盖汇总", 3)
    status_table(doc, [
        ["上行DL/T 645", "载波/双模/RS485，支持645及扩展", "RS485-I、载波、蓝牙共享645服务", "已实现主体"],
        ["下行有线通信", "RS485-II + 2×RJ45，每RJ45两路RS485", "五路逻辑下行串口，其中三路主采集、两路转接", "已实现主体"],
        ["下行无线", "最多12台，HRF，约200 m", "未见无线扫描、抄读和档案端口实现", "规范要求"],
        ["自动识别", "自动识别逆变器厂家/协议", "三主端口、地址1～10、特征范围匹配", "已实现，有范围限制"],
        ["协议转换", "主站统一645、下行厂家Modbus", "类型/字节序/倍率/BCD转换链完整", "已实现"],
        ["柔性调节", "数值、百分比、时段、广播", "启停、P/Q/PF、百分比、两类有功时段、FF全选", "已实现主体"],
        ["链路监测", "下行链路状态与异常记录", "超时/异常/数据失效存在，独立状态/事件未闭环", "部分实现"],
        ["通信转接", "通道切换/透明转发", "两组副口借用主口转发，单缓存协调", "已实现主体"],
        ["本地维护", "RS485-I/蓝牙，身份认证", "645维护通道存在，认证/权限未见", "部分实现"],
        ["事件记录", "9类事件及50/100条容量", "通用环形框架存在，业务只接入少量系统事件", "部分实现"],
        ["存储", "≥8 MB、分类、FIFO、升级不影响历史", "外部Flash分区、A/B和环形保存存在", "已实现主体，需容量验证"],
        ["本机电压", "单相，1.0级，误差±2%", "ADC/DMA RMS、校准并接入645", "已实现链路，精度需型式验证"],
        ["时钟", "日误差≤0.5 s、对时误差<10 s", "RTC和广播校时存在", "功能已接入，精度/事件待验证"],
        ["断点升级", "本地/远程、鉴权、断点续传", "位图/CRC/Boot搬运存在；串口应答和鉴权不足", "部分实现"],
    ], font_size=8.2)
    heading(doc, "16.2 点表与代码不一致项", 3)
    add_table(doc, ["项目", "文档/规范", "当前代码", "建议"], [
        ["无线串口参数", "点表含04000807", "数据中心仅04000800～06，注释明确不支持07", "补实现或从交付点表删除/标注不适用"],
        ["波特率上限", "规范支持最高115200", "校验表仅到38400", "扩展驱动和参数校验并回归测试"],
        ["协议库容量", "点表描述可到FE/约253槽", "代码容量100槽，645槽位按1～100", "统一点表、协议选择器与存储容量"],
        ["协议库写持久化", "维护后应长期有效", "645写只更新RAM，不保存Flash", "写成功后执行A/B持久化并记录变更事件"],
        ["默认抄读周期", "代码注释称默认10 s", "实际赋值0，表示关闭", "确定产品默认值并统一注释/测试"],
        ["串口参数生效", "维护常期待即时生效", "保存后需重启", "文档明确或增加安全重初始化"],
        ["升级响应通道", "本地/远程均应可升级", "升级应答固定走蓝牙HJ-02C", "升级消息保留来源并按原通道应答"],
        ["事件点", "规范9类事件", "当前仅重启、失电、上电、超级电容故障", "逐业务挂接并实现645事件读取/主动上报"],
    ], widths=[3.5, 4.6, 4.8, 4.0], font_size=8.8)
    heading(doc, "16.3 建议的闭环优先级", 3)
    add_table(doc, ["优先级", "闭环项", "原因"], [
        ["P0", "升级来源通道、升级鉴权和镜像真实性", "关系远程可用性及软件供应链安全"],
        ["P0", "本地维护身份认证与权限控制", "避免通过蓝牙/RS485直接改参、控制或升级"],
        ["P1", "规范事件全量接入及645查询/主动上报", "满足运维追溯和技术规范验收"],
        ["P1", "协议库容量、槽位范围和持久化统一", "避免维护后丢失或主站/设备理解不一致"],
        ["P1", "115200及04000807/无线功能决策", "消除接口规范与实际能力偏差"],
        ["P1", "链路监测状态与通信异常统计", "为故障定位和主动上报提供基础"],
        ["P2", "默认抄读周期、串口即时生效策略", "减少现场配置误解"],
        ["P2", "电压精度、RTC精度和恢复场景验证", "完成±2%和0.5 s/24 h指标证据"],
    ], widths=[2.0, 8.0, 6.9])

    heading(doc, "17 结论", 2, page_break=True)
    body_p(doc, "当前软件已经形成以“统一DL/T 645上行接口 + 协议库驱动的Modbus下行适配”为核心的完整骨架，主要业务包括自动识别、12台档案、周期/实时采集、启停与功率调节、两时段控制、通信转接、单相电压采集、参数和位置维护、事件与A/B存储、断点升级及Boot搬运。各模块通过档案和协议库解耦厂家差异，通过端口状态机和消息队列解决并发，通过有效性标志和FF应答避免错误数据上送。")
    body_p(doc, "交付前最需要闭环的是安全认证、升级原通道应答、规范事件全量接入、无线/04000807、115200波特率、协议库容量及持久化一致性。完成这些项目后，装置功能将从“主体链路可运行”进一步达到“规范、点表、代码、测试和现场运维一致”。")

    # The reference template carries automatic multilevel numbering on Heading 2-4.
    # Strip author-supplied numeric prefixes to avoid duplicated numbers in headings/TOC.
    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Heading 2", "Heading 3", "Heading 4"}:
            clean = re.sub(r"^\d+(?:\.\d+)*\s+", "", paragraph.text)
            if clean != paragraph.text:
                paragraph.text = clean

    # Footer and field updates.
    for section in doc.sections:
        footer = section.footer
        if footer.paragraphs:
            fp = footer.paragraphs[0]
            fp.clear()
        else:
            fp = footer.add_paragraph()
        add_page_field(fp)
    set_update_fields(doc)
    doc.core_properties.title = "DFDL33-SH8619 协议转换单元软件功能设计"
    doc.core_properties.subject = "协议转换单元总体框架、详细功能点及关联关系"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "依据0923技术规范、DLT64点表、接口资料及当前APP/Boot源码整理。"
    doc.save(TARGET)


if __name__ == "__main__":
    build()
